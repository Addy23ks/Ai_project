import html
from pathlib import Path
from typing import List, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable

from app.schemas.document import ValidationStatus, ConfidenceBreakdown
from app.schemas.model_a import UncertainSpan


class DocumentGeneratorService:
    def __init__(self):
        pass

    def generate_html_preview(
        self,
        transcription: str,
        uncertain_spans: List[UncertainSpan],
        status: ValidationStatus,
        score: float
    ) -> str:
        """
        Creates a clean, styled HTML preview of the typed document with highlighted unclear spans.
        """
        escaped_text = html.escape(transcription)

        # Highlight uncertain spans
        for span in uncertain_spans:
            if span.text:
                escaped_span = html.escape(span.text)
                replacement = (
                    f'<span class="bg-amber-100 border-b-2 border-amber-500 text-amber-900 px-1 py-0.5 rounded font-mono text-sm cursor-help" '
                    f'title="Unclear handwriting on page {span.page}: {html.escape(span.reason)}">'
                    f'{escaped_span} <span class="text-xs uppercase bg-amber-200 text-amber-800 px-1 rounded font-sans font-semibold">[UNCLEAR]</span></span>'
                )
                escaped_text = escaped_text.replace(escaped_span, replacement)

        # Convert linebreaks and paragraph breaks to HTML
        paragraphs = escaped_text.split("\n\n")
        html_paragraphs = []
        for p in paragraphs:
            p_clean = p.strip()
            if not p_clean:
                continue
            # Check if line looks like a header (e.g. uppercase or title)
            lines = p_clean.split("\n")
            if len(lines) == 1 and len(lines[0]) < 60 and (lines[0].isupper() or lines[0].endswith(":")):
                html_paragraphs.append(f'<h3 class="text-lg font-bold text-slate-800 mt-4 mb-2">{lines[0]}</h3>')
            elif p_clean.startswith("1.") or p_clean.startswith("- ") or p_clean.startswith("* "):
                # List items
                list_items = "".join([f'<li class="ml-4 mb-1">{line}</li>' for line in lines if line.strip()])
                html_paragraphs.append(f'<ul class="list-disc list-inside space-y-1 my-2 text-slate-700">{list_items}</ul>')
            else:
                formatted_p = "<br />".join(lines)
                html_paragraphs.append(f'<p class="mb-3 text-slate-700 leading-relaxed">{formatted_p}</p>')

        return "\n".join(html_paragraphs)

    def generate_docx(
        self,
        transcription: str,
        uncertain_spans: List[UncertainSpan],
        status: ValidationStatus,
        score: float,
        output_path: Path
    ) -> Path:
        """
        Generates a professionally formatted Word document (.docx).
        """
        doc = docx.Document()

        # Document Title / Metadata Header
        title = doc.add_heading("Extracted Handwritten Document", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # Summary Header
        p_meta = doc.add_paragraph()
        p_meta.paragraph_format.space_after = Pt(14)
        run_meta = p_meta.add_run(
            f"Extraction Confidence: {score}% | Status: {status.value} | Flagged Segments: {len(uncertain_spans)}"
        )
        run_meta.font.size = Pt(9.5)
        run_meta.font.color.rgb = RGBColor(100, 116, 139)  # Slate 500

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # Content parsing
        paragraphs = transcription.split("\n\n")
        for p in paragraphs:
            lines = p.strip().split("\n")
            if not lines or not lines[0].strip():
                continue

            # Header detection
            if len(lines) == 1 and len(lines[0]) < 60 and (lines[0].isupper() or lines[0].endswith(":")):
                h = doc.add_heading(lines[0], level=2)
                h.paragraph_format.space_before = Pt(12)
                h.paragraph_format.space_after = Pt(4)
                continue

            # Standard paragraph with potential uncertain spans
            p_doc = doc.add_paragraph()
            p_doc.paragraph_format.space_after = Pt(8)
            p_doc.paragraph_format.line_spacing = 1.15

            full_line_text = "\n".join(lines)

            # Check if any uncertain spans exist in this block
            matched_spans = [s for s in uncertain_spans if s.text and s.text in full_line_text]
            if matched_spans:
                # Add with explicit unclear indicator
                p_doc.add_run(full_line_text)
                for s in matched_spans:
                    p_flag = doc.add_paragraph()
                    r_flag = p_flag.add_run(f"[UNCLEAR — REVIEW REQUIRED: '{s.text}' (Page {s.page}: {s.reason})]")
                    r_flag.font.size = Pt(8.5)
                    r_flag.font.color.rgb = RGBColor(217, 119, 6)  # Amber
                    r_flag.font.italic = True
            else:
                p_doc.add_run(full_line_text)

        doc.save(str(output_path))
        return output_path

    def generate_pdf(
        self,
        transcription: str,
        uncertain_spans: List[UncertainSpan],
        status: ValidationStatus,
        score: float,
        output_path: Path
    ) -> Path:
        """
        Generates a clean PDF document using ReportLab.
        """
        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54
        )

        styles = getSampleStyleSheet()
        normal_style = styles["Normal"]
        normal_style.fontSize = 10
        normal_style.leading = 14
        normal_style.textColor = colors.HexColor("#334155")

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#0f172a"),
            spaceAfter=6
        )

        meta_style = ParagraphStyle(
            "DocMeta",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#64748b"),
            spaceAfter=12
        )

        h2_style = ParagraphStyle(
            "DocH2",
            parent=styles["Heading2"],
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#1e293b"),
            spaceBefore=10,
            spaceAfter=4
        )

        unclear_style = ParagraphStyle(
            "UnclearNote",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#d97706"),
            spaceAfter=4
        )

        story = []

        # Header
        story.append(Paragraph("Extracted Handwritten Document", title_style))
        story.append(Paragraph(
            f"Extraction Confidence: <b>{score}%</b> &nbsp;|&nbsp; Status: <b>{status.value}</b> &nbsp;|&nbsp; Flagged Segments: {len(uncertain_spans)}",
            meta_style
        ))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#e2e8f0"), spaceAfter=14))

        # Content
        paragraphs = transcription.split("\n\n")
        for p in paragraphs:
            lines = [l.strip() for l in p.strip().split("\n") if l.strip()]
            if not lines:
                continue

            if len(lines) == 1 and len(lines[0]) < 60 and (lines[0].isupper() or lines[0].endswith(":")):
                story.append(Paragraph(html.escape(lines[0]), h2_style))
                continue

            # Format lines
            text_block = "<br/>".join([html.escape(l) for l in lines])

            # Highlight unclear spans
            for span in uncertain_spans:
                if span.text and html.escape(span.text) in text_block:
                    escaped_s = html.escape(span.text)
                    text_block = text_block.replace(
                        escaped_s,
                        f'<font color="#b45309"><b>{escaped_s} [UNCLEAR]</b></font>'
                    )

            story.append(Paragraph(text_block, normal_style))
            story.append(Spacer(1, 8))

        doc.build(story)
        return output_path


document_generator = DocumentGeneratorService()
